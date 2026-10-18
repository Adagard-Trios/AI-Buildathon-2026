"""Fill the official AI Builder Cup submission template with FirstStep's content.

Run with python-pptx:  python build_deck.py  ->  FirstStep_AI_Builder_Cup_2026.pptx
Keeps the template's header strip, logos and slide headings; adds native, editable content.
"""

from copy import deepcopy

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = "source-template.pptx"
OUT = "FirstStep_AI_Builder_Cup_2026.pptx"
FONT = "Arial"

INK = RGBColor(0x14, 0x20, 0x1F)
MUTED = RGBColor(0x55, 0x62, 0x5F)
ACCENT = RGBColor(0x0B, 0x6E, 0x4F)
TINT = RGBColor(0xE7, 0xF4, 0xEE)
DANGER = RGBColor(0xB4, 0x23, 0x18)
DANGER_TINT = RGBColor(0xFD, 0xEC, 0xEA)
CAUTION = RGBColor(0x8A, 0x53, 0x00)
CAUTION_TINT = RGBColor(0xFD, 0xF3, 0xE1)
SUNKEN = RGBColor(0xF3, 0xF5, 0xF4)
LINE = RGBColor(0xDD, 0xE3, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LIVE_URL = "https://firststep-web-214625202650.us-central1.run.app"
REPO_URL = "https://github.com/Adagard-Trios/AI-Buildathon-2026"

prs = Presentation(SRC)
SLIDES = list(prs.slides)          # template order, 1-based below via S(n)


def S(n):
    return SLIDES[n - 1]


# ---------------------------------------------------------------- helpers

def text(slide, x, y, w, h, runs, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, spacing=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        pieces = para if isinstance(para, list) else [(para, {})]
        for piece, opts in pieces:
            r = p.add_run()
            r.text = piece
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = opts.get("color", color)
    return box


def box(slide, x, y, w, h, fill=SUNKEN, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, title, body, fill=SUNKEN, title_color=INK, size=11, title_size=13, chip=None):
    box(slide, x, y, w, h, fill=fill)
    pad = 0.14
    ty = y + pad
    if chip:
        c = box(slide, x + pad, ty, 0.34, 0.34, fill=ACCENT, shape=MSO_SHAPE.OVAL)
        text(slide, x + pad, ty, 0.34, 0.34, chip, size=11, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + pad + 0.44, ty, w - 2 * pad - 0.44, 0.34, title, size=title_size, bold=True,
             color=title_color, anchor=MSO_ANCHOR.MIDDLE)
        ty += 0.44
    else:
        text(slide, x + pad, ty, w - 2 * pad, 0.3, title, size=title_size, bold=True, color=title_color)
        ty += 0.32
    text(slide, x + pad, ty, w - 2 * pad, h - (ty - y) - pad, body, size=size, color=MUTED, spacing=3)


def arrow(slide, x, y, w=0.22, h=0.22, color=ACCENT):
    a = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def picture(slide, path, x, y, w=None, h=None, crop=None):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if crop:
        ratio = pic.height / pic.width
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        # keep the aspect ratio of what remains, or the crop is stretched to the old frame
        pic.height = int(pic.width * ratio * (1 - crop[1] - crop[3]) / (1 - crop[0] - crop[2]))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return pic


def heading(slide, new_text):
    """Rename the template's heading text box in place (keeps its font and position)."""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            run = sh.text_frame.paragraphs[0].runs[0]
            run.text = new_text
            for extra in sh.text_frame.paragraphs[0].runs[1:]:
                extra.text = ""
            # drop the template's helper paragraphs (questions, bullets) below the heading
            for p in list(sh.text_frame.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            run.font.size = Pt(18)
            sh.height = Inches(0.45)
            return sh
    raise RuntimeError("no heading on slide")


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ---------------------------------------------------------------- 1. team details
s = S(1)
tb = next(sh for sh in s.shapes if sh.has_text_frame and "Team Details" in sh.text_frame.text)
values = {"Team name:": "Adagard",
          "Team leader name:": "Nivakaran Shanmugabavan",
          "Problem Statement:": "Challenge 1 (BFSI) · Stopping the payment a scam victim approves (APP fraud)"}
for p in tb.text_frame.paragraphs:
    for label, value in values.items():
        if p.runs and p.runs[0].text.strip() == label:
            r = p.add_run()
            r.text = value
            r.font.name = FONT
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = ACCENT

notes(s, "FirstStep: an AI copilot for first-time investors in Sri Lanka, built on Gemini and Google Cloud.")

# ---------------------------------------------------------------- 3. brief
s = S(3)
heading(s, "Brief about the idea")
text(s, 0.3, 1.22, 9.4, 0.62,
     [[("FirstStep is the last checkpoint before the money leaves. ", {"size": 13, "bold": True, "color": ACCENT}),
       ("It catches the scam at the moment of payment, explains it to the customer in Sinhala, Tamil or "
        "English, and gives the bank and regulator a record that can't be quietly changed.", {"size": 13})]],
     name="Idea")
text(s, 0.3, 2.02, 9.4, 0.45,
     [[("In an authorised push payment scam the customer passes every check herself: right login, right OTP, "
        "her own decision to pay. Regulators are moving that loss onto banks.", {"size": 10.5, "color": MUTED})]])
stats = [("Oct 2024", "UK: mandatory reimbursement of APP scam losses, up to GBP 85,000 per claim"),
         ("Dec 2024", "Singapore: Shared Responsibility Framework puts duties on banks and telcos"),
         ("Jun 2026", "Sri Lanka: CBSL makes \u201cBe Scam Proof\u201d its consumer-protection theme")]
for i, (big, small) in enumerate(stats):
    x = 0.3 + i * 3.16
    box(s, x, 2.52, 3.0, 1.0, fill=TINT)
    text(s, x + 0.15, 2.58, 2.7, 0.4, big, size=20, bold=True, color=ACCENT)
    text(s, x + 0.15, 2.98, 2.7, 0.5, small, size=10, color=MUTED)
steps = [("1", "Pay a new payee", "Rs 30,000 to a new payee: before any code, the bank asks what it is for"),
         ("2", "Scam check", "She pastes the WhatsApp offer: STOP, and why, at the moment of payment"),
         ("3", "Explained", "In Sinhala, with every figure checked against the tools"),
         ("4", "Sealed record", "Evidence for disputes, reimbursement and the regulator")]
for i, (n, t, b) in enumerate(steps):
    x = 0.3 + i * 2.37
    card(s, x, 3.7, 2.2, 1.4, t, b, chip=n, size=10.5)
    if i < 3:
        arrow(s, x + 2.2 - 0.02, 4.3, 0.18, 0.2)
text(s, 0.3, 5.2, 9.4, 0.25,
     "Also: USD 88\u2013114 bn lost to online scams across Asia-Pacific in 2025 (UNODC); only 22% of victims recover anything (GASA).",
     size=9, color=MUTED)
notes(s, "Sources: UK PSR APP reimbursement (7 Oct 2024); Singapore MAS/IMDA Shared Responsibility Framework (16 Dec 2024); "
         "Newsfirst 2 Jun 2026 (CBSL 'Be Scam Proof'); UNODC 2025; GASA 2025.")

# ---------------------------------------------------------------- 4. opportunities
s = S(4)
heading(s, "Opportunities")
cols = [("Why not existing apps?", "",
         ["Google Messages and ScamShield protect the phone and need the victim to doubt the offer: 53% of victims aged 65+ who didn't check thought it unnecessary (CommBank, 2026)",
          "FirstStep protects the payment: it acts at the transfer, on any phone, whatever the channel (WhatsApp, calls)",
          "A bank-side control, so it fits Singapore's and the UK's duties on banks"]),
        ("How it solves the problem", "",
         ["At a new payee the bank pauses and asks what the payment is for, the moment regulators single out as high risk",
          "Plain-language reasons in the customer's language, where generic English warnings fail",
          "Policy decides, Gemini reads: the bank can tell a regulator exactly why a payment was stopped",
          "A sealed, verifiable record for disputes and reimbursement claims"]),
        ("USP", "",
         ["Protects the payment, not just the phone",
          "Live and tested on Cloud Run: 11/11 end-to-end checks",
          "Built for Sri Lanka first, APAC-ready: Sinhala, Tamil, LKR and CSE today; a new market is mostly data (regulator register, scam corpus, OKF pattern pack, Gemini language) plus a new policy version, no new services",
          "Next markets: Singapore (bank duties under the Shared Responsibility Framework); Malaysia and India (Tamil already supported)"])]
for i, (title, _, items) in enumerate(cols):
    x = 0.3 + i * 3.16
    box(s, x, 1.3, 3.0, 3.6, fill=TINT if i == 2 else SUNKEN)
    text(s, x + 0.18, 1.45, 2.65, 0.35, title, size=14, bold=True, color=ACCENT if i == 2 else INK)
    text(s, x + 0.18, 1.9, 2.65, 3.35, [[("•  ", {"color": ACCENT, "bold": True}), (it, {})] for it in items],
         size=10.5, color=INK, spacing=5)

# ---------------------------------------------------------------- 5. features
s = S(5)
heading(s, "List of features offered by the solution")
feats = [("Payment checkpoint", "A new payee or unusual behaviour pauses the transfer: the bank asks what it is for, checks the answer with Scam Shield, and only then sends a one-time code."),
         ("Scam Shield", "Paste a WhatsApp, SMS or email offer. Gemini reads the red flags, policy SHIELD-P001 decides: a STOP / CHECK / CLEAR stamp with what to do next."),
         ("Explained in her language", "Ask why, by text or voice, in Sinhala, Tamil or English. Gemini answers from the platform's tools; any unchecked figure is withheld."),
         ("Sealed evidence", "Each decision is hashed with its model and policy and anchored in a contract; change one field and verification fails. Readable over OData."),
         ("Scam knowledge, open", "Scam patterns in Google's Open Knowledge Format, with sources and trust tiers, cited in every verdict."),
         ("Slowing down is cheap", "Scammers sell urgency. The cash planner shows what rushing a sale costs, and the cheapest way to raise money calmly.")]
for i, (t, b) in enumerate(feats):
    x = 0.3 + (i % 3) * 3.16
    y = 1.3 + (i // 3) * 1.85
    card(s, x, y, 3.0, 1.7, t, b, chip=str(i + 1), size=10.5)

# ---------------------------------------------------------------- 6. process flow
s = S(6)
heading(s, "Process flow diagram")
flow = [("Transfer", "Kamala pays Rs 30,000 to a new payee", SUNKEN, INK),
        ("Gateway", "Behavioural and payee-network risk score (heuristic)", SUNKEN, INK),
        ("Pause & ask", "\u201cWhat is this payment for?\u201d She pastes the WhatsApp offer", CAUTION_TINT, CAUTION),
        ("Scam check", "Gemini reads the red flags; SHIELD-P001 decides", TINT, ACCENT),
        ("STOP", "Cancel is the default; continuing needs consent, then the code", DANGER_TINT, DANGER)]
w = 1.72
for i, (t, b, fill, col) in enumerate(flow):
    x = 0.3 + i * (w + 0.2)
    box(s, x, 1.35, w, 1.45, fill=fill)
    text(s, x + 0.12, 1.45, w - 0.24, 0.32, t, size=13, bold=True, color=col)
    text(s, x + 0.12, 1.8, w - 0.24, 0.95, b, size=10, color=MUTED)
    if i < len(flow) - 1:
        arrow(s, x + w + 0.02, 1.98, 0.16, 0.2)
text(s, 0.3, 2.95, 9.4, 0.3, "Then the customer's question and the bank's record", size=12, bold=True, color=INK)
flow2 = [("Ask why", "Assistant answers in Sinhala; invented figures are withheld", TINT, ACCENT),
         ("Decide", "She cancels, or acknowledges the risk and confirms with a code", SUNKEN, INK),
         ("Sealed", "Verdict and transfer decision hashed with model and policy", SUNKEN, INK),
         ("Verify", "Any change later fails against the contract", SUNKEN, INK),
         ("Report", "Bank and regulator read every decision over OData", SUNKEN, INK)]
for i, (t, b, fill, col) in enumerate(flow2):
    x = 0.3 + i * (w + 0.2)
    box(s, x, 3.35, w, 1.45, fill=fill)
    text(s, x + 0.12, 3.45, w - 0.24, 0.32, t, size=13, bold=True, color=col)
    text(s, x + 0.12, 3.8, w - 0.24, 0.95, b, size=10, color=MUTED)
    if i < len(flow2) - 1:
        arrow(s, x + w + 0.02, 3.98, 0.16, 0.2)
text(s, 0.3, 4.95, 9.4, 0.4,
     "The step-up runs Scam Shield on the customer's answer before any code is sent; feeding the verdict into the gateway's risk score is next.",
     size=10.5, color=MUTED)

# ---------------------------------------------------------------- 7. wireframes / design
s = S(7)
heading(s, "Design: the verdict is the product")
text(s, 0.3, 1.25, 4.3, 1.4,
     [[("One decision, one stamp.", {"size": 14, "bold": True})],
      [("The visual world is a bank passbook: paper, ink, savings-book green for safe, stamp red for "
        "stop. The verdict is a stamp with a word and a shape, never colour alone, then what to do, "
        "then why.", {"size": 11, "color": MUTED})]], spacing=6)
stamps = [("STOP", "This looks like a scam", DANGER, DANGER_TINT),
          ("CHECK", "Some warning signs", CAUTION, CAUTION_TINT),
          ("CLEAR", "No warning signs found", ACCENT, TINT),
          ("UNSURE", "Too little to call it safe", MUTED, SUNKEN)]
for i, (word, line, col, tint) in enumerate(stamps):
    x = 0.3 + (i % 2) * 2.2
    y = 2.85 + (i // 2) * 1.25
    box(s, x, y, 2.05, 1.1, fill=tint)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.12), Inches(y + 0.15), Inches(0.8), Inches(0.8))
    o.fill.background()
    o.line.color.rgb = col
    o.line.width = Pt(2.25)
    o.shadow.inherit = False
    text(s, x + 0.12, y + 0.15, 0.8, 0.8, word, size=9 if len(word) > 5 else 10, bold=True, color=col,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.0, y + 0.2, 0.98, 0.75, line, size=10, color=INK, anchor=MSO_ANCHOR.MIDDLE)
picture(s, "shot-verdict.png", 4.85, 1.25, w=4.85)
text(s, 4.85, 4.38, 4.85, 0.9,
     "Live screen. Every text colour clears 4.5:1 contrast in light and dark; layouts are checked at phone width (390 px).",
     size=10, color=MUTED)

# ---------------------------------------------------------------- 8. architecture
s = S(8)
heading(s, "Architecture diagram of the proposed solution")
# users
box(s, 0.3, 1.3, 1.55, 0.75, fill=SUNKEN)
text(s, 0.4, 1.36, 1.35, 0.65, [[("Customers", {"bold": True, "size": 11})], [("Web app and mobile app", {"size": 9, "color": MUTED})]])
box(s, 0.3, 2.2, 1.55, 0.75, fill=SUNKEN)
text(s, 0.4, 2.26, 1.35, 0.65, [[("Banks, regulators", {"bold": True, "size": 11})], [("OData v4 clients (Excel, BI)", {"size": 9, "color": MUTED})]])
arrow(s, 1.92, 1.97, 0.16, 0.22)
# cloud run frame
frame = box(s, 2.2, 1.25, 5.45, 3.95, fill=None, line=ACCENT, radius=0.04)
text(s, 2.35, 1.3, 5.2, 0.28, "Google Cloud Run  ·  each service verifies every request itself",
     size=9.5, bold=True, color=ACCENT)
svcs = [("Web (Next.js)", "Scam Shield, planner, audit, assistant UI"),
        ("Platform", "Accounts, roles, portfolios, token signing"),
        ("C2 Fraud & Scam", "Gemini red flags, SHIELD-P001, gateway G1-G7, step-up, OKF bundle"),
        ("C1 Portfolio", "Fuzzy GA cash planner, MOEA/D, forecast bands, statement import"),
        ("C4 Assistant", "LangGraph agent on Gemini, grounding check, Trust Panel"),
        ("C3 Audit", "Hash-committed decision log; AuditRegistry.sol on the audit service chain")]
for i, (t, b) in enumerate(svcs):
    x = 2.35 + (i % 2) * 2.62
    y = 1.68 + (i // 2) * 1.15
    box(s, x, y, 2.5, 1.02, fill=TINT if i in (2, 4) else SUNKEN)
    text(s, x + 0.1, y + 0.07, 2.3, 0.3, t, size=11, bold=True)
    text(s, x + 0.1, y + 0.37, 2.3, 0.62, b, size=9, color=MUTED)
# google services on the right
right = [("Gemini", "Red flags, agent reasoning, Sinhala and Tamil answers"),
         ("Secret Manager", "Signing key, service-client secrets"),
         ("Cloud Build", "Images per service, kept in Artifact Registry"),
         ("GitHub Actions + WIF", "Keyless deploy on every push")]
for i, (t, b) in enumerate(right):
    y = 1.25 + i * 1.0
    box(s, 7.85, y, 1.85, 0.88, fill=TINT if i == 0 else SUNKEN)
    text(s, 7.95, y + 0.06, 1.65, 0.3, t, size=10.5, bold=True, color=ACCENT if i == 0 else INK)
    text(s, 7.95, y + 0.36, 1.65, 0.5, b, size=8.5, color=MUTED)
arrow(s, 7.67, 1.55, 0.14, 0.2)

# ---------------------------------------------------------------- 9. technologies
s = S(9)
heading(s, "Technologies used in the solution")
rows = [("Live", "Gemini (red flags, agent, Sinhala/Tamil) · Cloud Run (6 services) · Cloud Build · Artifact Registry · "
                 "Secret Manager · Workload Identity Federation · Open Knowledge Format (OKF) · Cloud Logging"),
        ("Live", "Next.js 16 + React 19 · FastAPI · Express · LangGraph · Expo (React Native) · Hardhat EVM + Solidity · "
                 "OData v4 · RS256 / JWKS"),
        ("Offline", "TimesFM fine-tuned with LoRA on a Colab T4 + LSTM residual head; walk-forward evaluation"),
        ("Planned", "Vertex AI, ADK + Agent Engine, BigQuery ML AI.FORECAST, Document AI, Vertex AI Search, "
                    "Pub/Sub, Firestore, Firebase Auth, Cloud Armor, VPC-SC, Blockchain Node Engine")]
colors = {"Live": (TINT, ACCENT), "Offline": (CAUTION_TINT, CAUTION), "Planned": (SUNKEN, MUTED)}
for i, (status, body) in enumerate(rows):
    y = 1.3 + i * 1.0
    fill, col = colors[status]
    box(s, 0.3, y, 9.4, 0.88, fill=fill)
    pill = box(s, 0.45, y + 0.27, 1.0, 0.34, fill=col, radius=0.5)
    text(s, 0.45, y + 0.27, 1.0, 0.34, status, size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.65, y + 0.1, 7.9, 0.7, body, size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.3, 5.3, 9.4, 0.22, "Status as deployed on 18 Oct 2026. The README's technology map lists the same split.",
     size=9, color=MUTED)

# ---------------------------------------------------------------- 11. snapshots
s = S(11)
heading(s, "Snapshots of the prototype")
shots = [("shot-stepup.png", "Payment checkpoint: the bank asks what the payment is for and shows STOP before any code"),
         ("shot-verdict.png", "Scam Shield: STOP stamp, reasons, cited OKF scam patterns; red flags read by Gemini"),
         ("shot-assistant.png", "Assistant answering in Sinhala, from the scam-check tool"),
         ("shot-tamper.png", "Audit: record altered, verification fails against the contract")]
for i, (f, cap) in enumerate(shots):
    x = 0.3 + (i % 2) * 4.75
    y = 1.25 + (i // 2) * 2.1
    picture(s, f, x, y, w=3.05, crop=(0, 0.125, 0, 0))
    text(s, x + 3.15, y + 0.1, 1.45, 1.7, cap, size=10, color=MUTED)

# ---------------------------------------------------------------- 12. benchmarking
s = S(12)
heading(s, "Prototype performance report / benchmarking")
cd = CategoryChartData()
cd.categories = ["Rules only", "Gemini + rules (live)"]
cd.add_series("Scams caught (recall)", (0.45, 1.0))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.3), Inches(1.3), Inches(4.6), Inches(2.1), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Scam Shield: scams caught, 40 labelled messages"
for p in ch.chart_title.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.name = FONT
plot = ch.plots[0]
plot.has_data_labels = True
plot.data_labels.number_format = "0%"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(10)
plot.data_labels.font.name = FONT
plot.gap_width = 60
pts = plot.series[0].points
for idx, col in ((0, MUTED), (1, ACCENT)):
    pts[idx].format.fill.solid()
    pts[idx].format.fill.fore_color.rgb = col
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = 1.25
ch.category_axis.tick_labels.font.size = Pt(10)
ch.category_axis.tick_labels.font.name = FONT
text(s, 0.3, 3.5, 4.6, 1.85,
     [[("Synthetic test set (40 msgs): recall 45% \u2192 100% with Gemini, 0 false positives", {"size": 13, "bold": True, "color": ACCENT})],
      [("Gemini reading red flags lifts recall from 45% to 100% at the same zero false-positive rate "
        "(precision 1.00). Honest limits: 40 synthetic English messages written by the team "
        "(eval/scam_synthetic_v1.jsonl); Sinhala and Tamil not yet measured; a real labelled "
        "Sri Lankan set is next.", {"size": 9.5, "color": MUTED})],
      [("Forecasting (offline): fine-tuned TimesFM hybrid, \u22126.8% pinball loss vs LSTM.", {"size": 9.5, "color": MUTED})]],
     spacing=4)
checks = [("11 / 11", "live end-to-end checks pass on Cloud Run"),
          ("176 +", "automated tests passing: platform 63, fraud 46, assistant 22, web 38, mobile 7, plus the audit suite"),
          ("401 / 403", "no token or wrong role is refused, by every service"),
          ("6 / 6", "services live on Cloud Run, redeployed by CI on every push")]
for i, (big, small) in enumerate(checks):
    y = 1.3 + i * 1.0
    box(s, 5.15, y, 4.55, 0.88, fill=SUNKEN)
    text(s, 5.3, y + 0.12, 1.35, 0.64, big, size=17, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.7, y + 0.1, 2.9, 0.7, small, size=10, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 13. future
s = S(13)
heading(s, "Additional details / future development")
done_items = [("Honest by design", "The README separates what is live, offline and planned; negative results are reported."),
              ("Responsible AI", "Policies decide, figures are grounded, privacy and advice checks run on every answer; family guardian alerts wait for the user's consent.")]
next_items = [("Verdict into the score", "Feed the scam verdict into the gateway's risk score, not only the step-up."),
              ("Real-world eval", "Labelled Sri Lankan messages, Sinhala and Tamil, with partner banks."),
              ("Gemini on Vertex AI + ADK", "Move the agent to Agent Engine; Document AI for statements."),
              ("Public anchoring", "Anchor the log on a public testnet via Blockchain Node Engine."),
              ("WhatsApp bot", "Forward a message, get the verdict back in the same chat.")]
text(s, 0.3, 1.25, 4.4, 0.3, "Already in place", size=13, bold=True, color=ACCENT)
for i, (t, b) in enumerate(done_items):
    card(s, 0.3, 1.65 + i * 1.75, 4.4, 1.6, t, b, fill=TINT, size=11)
text(s, 5.0, 1.25, 4.7, 0.3, "Next", size=13, bold=True, color=INK)
for i, (t, b) in enumerate(next_items):
    y = 1.65 + i * 0.73
    box(s, 5.0, y, 4.7, 0.63, fill=SUNKEN)
    text(s, 5.12, y + 0.06, 1.75, 0.52, t, size=10.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.95, y + 0.06, 2.65, 0.52, b, size=9.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 14. links
s = S(14)
tbx = next(sh for sh in s.shapes if sh.has_text_frame and "Provide links" in sh.text_frame.text)
links = {"GitHub Public Repository": REPO_URL,
         "Demo Video Link (3 Minutes)": "ADD THE VIDEO LINK",
         "Final Product Link": LIVE_URL}
for p in tbx.text_frame.paragraphs:
    for label, url in links.items():
        if p.runs and p.runs[0].text.strip() == label:
            p.runs[0].text = label + ":  "
            r = p.add_run()
            r.text = url
            r.font.name = FONT
            r.font.size = Pt(12)
            r.font.color.rgb = DANGER if url.startswith("ADD") else ACCENT
            if url.startswith("http"):
                r.hyperlink.address = url
tbx.height = Inches(1.6)
box(s, 0.3, 2.65, 9.4, 2.2, fill=TINT)
text(s, 0.5, 2.8, 9.0, 1.95,
     [[("Try it", {"size": 14, "bold": True, "color": ACCENT})],
      [("Register any email: every account starts with a demo portfolio of CSE, US and crypto holdings.", {})],
      [("Fraud → Scam Shield: paste an investment offer and see the verdict, reasons and scam patterns.", {})],
      [("Fraud → Check a transfer: tick New beneficiary, amount 30,000 → STEP-UP → enter the code.", {})],
      [("Audit → Tamper-evidence: anchor, verify, tamper, verify again.", {})],
      [("Services scale to zero to keep cost near nothing: the first click on a quiet service takes ~15 s.",
        {"color": MUTED, "size": 10})]], size=11.5, spacing=5)

# ---------------------------------------------------------------- 15 (blank) -> zero trust & honesty
s = S(15)
hb = text(s, 0.39, 0.78, 9.3, 0.45, "Built to be trusted", size=18, color=RGBColor(0, 0, 0))
for r in hb.text_frame.paragraphs[0].runs:
    r.font.name = "Google Sans Flex Medium"
zt = [("Verify every request", "Each service checks RS256 tokens against the platform's JWKS; no token means 401, wrong audience or role 403."),
      ("Least privilege", "Roles come from deployment config only; staff addresses cannot be self-registered."),
      ("Your data only", "Row-level reads: customers see only their own records; auditors read all, alter nothing."),
      ("Fresh proof for risk", "STEP-UP one-time codes on risky transfers, even inside a valid session."),
      ("Fail closed", "Misconfiguration refuses instead of allowing; denials are sealed in the audit trail."),
      ("No invented numbers", "An answer with a figure no tool produced is withheld.")]
for i, (t, b) in enumerate(zt):
    x = 0.3 + (i % 3) * 3.16
    y = 1.3 + (i // 3) * 1.85
    card(s, x, y, 3.0, 1.7, t, b, chip=str(i + 1), size=10.5)

# ---------------------------------------------------------------- delete the instruction slide and the optional cost slide
def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sid in list(sld_ids):
        if prs.slides.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            sld_ids.remove(sid)
            return
    raise RuntimeError("slide not found")


delete_slide(prs, S(2))
delete_slide(prs, S(10))

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
